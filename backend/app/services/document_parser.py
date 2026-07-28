"""Document text extraction for PDF and DOCX files.

Everything is processed in-memory — no files are saved to disk.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

import fitz  # PyMuPDF
from docx import Document as DocxDocument

from app.core.config import settings


class ParseError(Exception):
    """Raised when a document cannot be parsed."""


@dataclass
class ParsedDocument:
    filename: str
    size_bytes: int
    file_type: str
    pages: int
    text: str


def _validate(filename: str, size: int) -> str:
    name = filename.lower()
    ext = ".pdf" if name.endswith(".pdf") else ".docx" if name.endswith(".docx") else ""
    if ext not in settings.allowed_extensions:
        raise ParseError("Unsupported file type. Please upload a PDF or DOCX file.")
    if size > settings.max_file_size_bytes:
        raise ParseError(
            f"File is too large ({size / 1024 / 1024:.1f} MB). "
            f"Maximum is {settings.max_file_size_bytes / 1024 / 1024:.0f} MB."
        )
    return ext.lstrip(".")


def parse_pdf(data: bytes, filename: str, size: int) -> ParsedDocument:
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParseError("Could not open this PDF. It may be corrupted or encrypted.") from exc

    if doc.needs_pass:
        doc.close()
        raise ParseError("This PDF is password-protected. Please remove the password and try again.")

    pages = doc.page_count
    text_parts: list[str] = []
    for page in doc:
        text_parts.append(page.get_text("text"))
    doc.close()

    text = "\n\n".join(text_parts).strip()
    if not text:
        raise ParseError("No readable text found — the PDF may be a scanned image without a text layer.")
    return ParsedDocument(filename, size, "pdf", pages, text)


def parse_docx(data: bytes, filename: str, size: int) -> ParsedDocument:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("Could not read this DOCX file. It may be corrupted.") from exc

    text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(text_parts).strip()
    if not text:
        raise ParseError("No readable text found in this DOCX file.")

    words = len(text.split())
    pages = max(1, -(-words // 500))  # ceil division
    return ParsedDocument(filename, size, "docx", pages, text)


def parse_document(data: bytes, filename: str, size: int) -> ParsedDocument:
    file_type = _validate(filename, size)
    if file_type == "pdf":
        return parse_pdf(data, filename, size)
    return parse_docx(data, filename, size)
